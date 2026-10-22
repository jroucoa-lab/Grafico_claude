#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generador_graficos.py  (v2)
===========================
Analisis tecnico diario + semanal con cache incremental de datos.

QUE ANALIZA
  - tickers.txt con tickers (comas, lineas o ';')  -> analiza solo esos
  - tickers.txt vacio (o inexistente)              -> analiza todo el S&P 500

DATOS (cache/)
  - 1a ejecucion: descarga el historico completo y lo guarda (un CSV por ticker).
  - Siguientes: solo descarga los dias que faltan. Nunca guarda barras parciales
    (si el mercado de EE. UU. no ha cerrado, descarta la barra de hoy).
  - Si un split/dividendo cambia el historico, redescarga ese ticker entero.
  - Precios: Close de Yahoo (ajustado por splits, NO por dividendos) = lo que
    ves en tu broker. Los niveles y el plan estan en precios reales.

SALIDAS (carpeta --salida)
  ranking.csv            ranking con entrada, stop, TP1/TP2, acciones, capital
  resumen_patrones.xlsx  hojas Ranking, Resumen, Velas, Figuras, Avisos
  fichas.md              ficha tecnica por ticker + ultimas barras OHLCV
  graficos_tecnicos.pdf  2 paginas por ticker (diario + semanal), ordenadas por ranking
  compras.txt            tickers en Compra, separados por comas (formato TradingView)
  estado_datos.csv       por ticker: primera/ultima barra, dias de atraso y estado
  avisos.txt             problemas de datos (huecos, sin barras nuevas, barra provisional)

USO
  python generador_graficos.py                         # lee tickers.txt
  python generador_graficos.py --tickers PFE,NUE,PR
  python generador_graficos.py --detalle completo      # limpio | normal | completo
  python generador_graficos.py --solo-cache            # sin red
  python generador_graficos.py --refrescar             # fuerza descarga completa
  python generador_graficos.py --telegram              # avisa si cambian las compras
  python generador_graficos.py --importar-csv datos_ohlcv.csv --descartar-ultima

Dependencias: pip install pandas numpy matplotlib yfinance openpyxl lxml

AVISO: detecciones heuristicas para REVISAR visualmente. No es asesoramiento financiero.
"""
import argparse
import io
import json
import math
import os
import re
import sys
import time
from datetime import datetime, timedelta, timezone

import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages

BENCH = "^GSPC"
COLS = ["Open", "High", "Low", "Close", "AdjClose", "Volume"]
SECTOR_ETF = {
    "Information Technology": "XLK", "Technology": "XLK",
    "Financials": "XLF", "Financial Services": "XLF",
    "Health Care": "XLV", "Healthcare": "XLV",
    "Energy": "XLE", "Industrials": "XLI",
    "Consumer Discretionary": "XLY", "Consumer Cyclical": "XLY",
    "Consumer Staples": "XLP", "Consumer Defensive": "XLP",
    "Materials": "XLB", "Basic Materials": "XLB",
    "Utilities": "XLU", "Real Estate": "XLRE",
    "Communication Services": "XLC",
}
ETFS = sorted(set(SECTOR_ETF.values()))
MIN_BARRAS_D = 220
MIN_BARRAS_W = 45
VERDE, ROJO = "#1a9850", "#d73027"


# ============================================================================
# Cache de datos
# ============================================================================
def ruta_cache(cache, t):
    return os.path.join(cache, t.replace("^", "_").replace("/", "_") + ".csv")


def leer_cache(cache, t):
    p = ruta_cache(cache, t)
    if not os.path.exists(p):
        return None
    try:
        df = pd.read_csv(p, parse_dates=["Fecha"], index_col="Fecha")
        return df[COLS].sort_index()
    except Exception:
        return None


def guardar_cache(cache, t, df):
    df = df[~df.index.duplicated(keep="last")].sort_index()
    df.index.name = "Fecha"
    df[COLS].to_csv(ruta_cache(cache, t))


def ahora_ny():
    try:
        from zoneinfo import ZoneInfo
        return datetime.now(ZoneInfo("America/New_York"))
    except Exception:
        return datetime.now(timezone.utc) - timedelta(hours=5)


def mercado_cerrado(ny):
    return ny.hour * 60 + ny.minute >= 16 * 60 + 30


def quitar_barra_parcial(df):
    """Descarta la barra de hoy si Wall Street aun no ha cerrado (cierre + 30 min)."""
    if df is None or df.empty:
        return df
    ny = ahora_ny()
    hoy = pd.Timestamp(ny.date())
    ult = df.index[-1]
    if ult > hoy or (ult == hoy and not mercado_cerrado(ny)):
        return df.iloc[:-1]
    return df


def ultimo_dia_esperado():
    ny = ahora_ny()
    d = pd.Timestamp(ny.date())
    if d.weekday() >= 5 or not mercado_cerrado(ny):
        d -= pd.Timedelta(days=1)
    while d.weekday() >= 5:
        d -= pd.Timedelta(days=1)
    return d


def _extraer(df, t):
    if df is None or df.empty:
        return None
    if isinstance(df.columns, pd.MultiIndex):
        lv0, lv1 = df.columns.get_level_values(0), df.columns.get_level_values(1)
        if t in lv0:
            x = df[t]
        elif t in lv1:
            x = df.xs(t, axis=1, level=1)
        else:
            return None
    else:
        x = df
    x = x.rename(columns={"Adj Close": "AdjClose"}).copy()
    if "AdjClose" not in x and "Close" in x:
        x["AdjClose"] = x["Close"]
    if any(c not in x for c in COLS):
        return None
    x = x[COLS].dropna(subset=["Close"])
    x["Volume"] = x["Volume"].fillna(0)
    if x.empty:
        return None
    x.index = pd.to_datetime(x.index).tz_localize(None).normalize()
    return x


def bajar_lote(tickers, start=None, años=6, reintentos=3):
    import yfinance as yf
    out, pend = {}, list(tickers)
    for intento in range(reintentos):
        if not pend:
            break
        fallidos = []
        for i in range(0, len(pend), 50):
            lote = pend[i:i + 50]
            df = None
            try:
                kw = dict(interval="1d", auto_adjust=False, progress=False,
                          group_by="ticker", threads=True)
                if start is not None:
                    kw["start"] = start.strftime("%Y-%m-%d")
                else:
                    kw["period"] = f"{años}y"
                df = yf.download(lote, **kw)
            except Exception as e:  # noqa
                print(f"  ! error en lote: {e}")
            for t in lote:
                x = _extraer(df, t)
                if x is None:
                    fallidos.append(t)
                else:
                    out[t] = x
        pend = fallidos
        if pend and intento < reintentos - 1:
            time.sleep(3 * (intento + 1))
    return out, pend


def consistente(viejo, nuevo, tol=5e-4):
    """Compara el solape: si Close/AdjClose cambian hubo split o dividendo."""
    com = viejo.index.intersection(nuevo.index)
    if len(com) == 0:
        return False
    a, b = viejo.loc[com], nuevo.loc[com]
    for c in ("Close", "AdjClose"):
        r = (b[c] / a[c]).replace([np.inf, -np.inf], np.nan).dropna()
        if len(r) and (r - 1).abs().max() > tol:
            return False
    return True


def actualizar_cache(tickers, cache, años, refrescar=False, solo_cache=False):
    os.makedirs(cache, exist_ok=True)
    datos, avisos = {}, []
    completos, ultimas = [], {}
    for t in tickers:
        c = None if refrescar else leer_cache(cache, t)
        if c is not None and len(c):
            datos[t] = c
            ultimas[t] = c.index[-1]
        else:
            completos.append(t)
    if solo_cache:
        avisos += [f"{t}: sin cache (--solo-cache)" for t in completos]
        return datos, avisos

    esperado = ultimo_dia_esperado()
    cortos = {}
    for t, u in ultimas.items():
        if u >= esperado:
            continue
        if (esperado - u).days <= 40:
            cortos[t] = u
        else:
            completos.append(t)
    if cortos:
        start = min(cortos.values()) - pd.Timedelta(days=7)
        print(f"Actualizando {len(cortos)} tickers desde {start:%Y-%m-%d}...")
        nuevos, fallidos = bajar_lote(list(cortos), start=start)
        for t, n in nuevos.items():
            n = quitar_barra_parcial(n)
            if n is None or n.empty:
                continue
            viejo = datos[t]
            if consistente(viejo, n):
                datos[t] = pd.concat([viejo[~viejo.index.isin(n.index)], n]).sort_index()
                guardar_cache(cache, t, datos[t])
            else:
                completos.append(t)
                avisos.append(f"{t}: cambio de ajuste (split/dividendo); se redescarga completo")
        for t in fallidos:
            avisos.append(f"{t}: no se pudo actualizar; cache hasta {datos[t].index[-1]:%Y-%m-%d}")
    completos = list(dict.fromkeys(completos))
    if completos:
        print(f"Descarga completa de {len(completos)} tickers ({años} años)...")
        nuevos, fallidos = bajar_lote(completos, años=años)
        for t, n in nuevos.items():
            n = quitar_barra_parcial(n)
            if n is not None and not n.empty:
                datos[t] = n
                guardar_cache(cache, t, n)
        for t in fallidos:
            avisos.append(f"{t}: sin datos de Yahoo" + (" (se usa cache antigua)" if t in datos else ""))
    return datos, avisos


def barra_provisional_por_volumen(datos):
    """Detecta la barra de hoy sin consolidar: Yahoo da volumenes que no son multiplos de 100
    hasta que cierra el dia. Se decide con el conjunto de tickers, no con uno solo."""
    ult = {t: d.index[-1] for t, d in datos.items() if not t.startswith("^") and len(d) > 3}
    if len(ult) < 5:
        return None, ""
    L = max(ult.values())
    en_L = [t for t, u in ult.items() if u == L]
    if len(en_L) < 5:
        return None, ""

    def frac(fecha):
        vs = [datos[t].loc[fecha, "Volume"] for t in en_L if fecha in datos[t].index]
        vs = [v for v in vs if v > 0]
        return (float(np.mean([v % 100 != 0 for v in vs])) if vs else np.nan), len(vs)

    prev = sorted({datos[t].index[-2] for t in en_L})[-1]
    fL, nL = frac(L)
    fP, _ = frac(prev)
    if nL >= 5 and np.isfinite(fL) and np.isfinite(fP) and fL >= 0.6 and fP <= 0.25:
        return L, (f"{fL * 100:.0f}% de los volumenes del {L:%d/%m} no son multiplos de 100 "
                   f"(el dia anterior {fP * 100:.0f}%)")
    return None, ""


def descartar_provisional(datos, cache):
    L, det = barra_provisional_por_volumen(datos)
    if L is None:
        return []
    n = 0
    for t, d in list(datos.items()):
        if len(d) and d.index[-1] == L:
            datos[t] = d.iloc[:-1]
            guardar_cache(cache, t, datos[t])
            n += 1
    return [f"Barra del {L:%Y-%m-%d} provisional ({det}): descartada en {n} series; "
            f"se recalcula en la proxima ejecucion"]


def importar_csv(ruta, cache, descartar_ultima=False):
    df = pd.read_csv(ruta, parse_dates=["Fecha"])
    if "AdjClose" not in df:
        df["AdjClose"] = df["Close"]
    if descartar_ultima:
        df = df[df["Fecha"] < df["Fecha"].max()]
    os.makedirs(cache, exist_ok=True)
    n = 0
    for t, g in df.groupby("Ticker"):
        nuevo = g.set_index("Fecha")[COLS].sort_index()
        viejo = leer_cache(cache, t)
        if viejo is not None:
            nuevo = pd.concat([viejo[~viejo.index.isin(nuevo.index)], nuevo]).sort_index()
        guardar_cache(cache, t, nuevo)
        n += 1
    print(f"Importados {n} tickers desde {ruta}")


# ============================================================================
# Universo
# ============================================================================
def parse_tickers(txt):
    lineas = [ln.split("#")[0] for ln in (txt or "").splitlines()]
    txt = ",".join(lineas).replace(";", ",")
    out, vistos = [], set()
    for t in re.split(r"[,\s]+", txt):
        t = t.strip().upper().replace(".", "-")
        if t and t not in vistos:
            vistos.add(t)
            out.append(t)
    return out


def cargar_sp500(cache):
    """Lista del S&P 500 (ticker, sector). Se refresca como mucho cada 30 dias."""
    ruta = os.path.join(cache, "sp500.csv")
    if os.path.exists(ruta) and time.time() - os.path.getmtime(ruta) < 30 * 86400:
        return pd.read_csv(ruta)
    try:
        import urllib.request
        req = urllib.request.Request(
            "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies",
            headers={"User-Agent": "Mozilla/5.0"})
        html = urllib.request.urlopen(req, timeout=30).read().decode("utf-8")
        tb = pd.read_html(io.StringIO(html))[0]
        df = pd.DataFrame({
            "Ticker": tb["Symbol"].astype(str).str.upper().str.replace(".", "-", regex=False),
            "Sector": tb["GICS Sector"]})
        if not 400 <= len(df) <= 600:
            raise ValueError(f"lista con {len(df)} filas")
        os.makedirs(cache, exist_ok=True)
        df.to_csv(ruta, index=False)
        return df
    except Exception as e:  # noqa
        print(f"Aviso: no se pudo leer la lista del S&P 500 ({e})")
        if os.path.exists(ruta):
            print("  usando la ultima copia guardada")
            return pd.read_csv(ruta)
        return None


# ============================================================================
# Indicadores
# ============================================================================
def sma(s, n):
    return s.rolling(n).mean()


def ema(s, n):
    return s.ewm(span=n, adjust=False, min_periods=n).mean()


def rsi(close, n=14):
    d = close.diff()
    up, dn = d.clip(lower=0), -d.clip(upper=0)
    au = up.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    ad = dn.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    return 100 - 100 / (1 + au / ad.replace(0, np.nan))


def macd(close, f=12, s=26, sg=9):
    m = close.ewm(span=f, adjust=False).mean() - close.ewm(span=s, adjust=False).mean()
    sig = m.ewm(span=sg, adjust=False).mean()
    return m, sig, m - sig


def estocastico(h, l, c, n=14, k=3, d=3):
    ll, hh = l.rolling(n).min(), h.rolling(n).max()
    raw = 100 * (c - ll) / (hh - ll).replace(0, np.nan)
    kk = raw.rolling(k).mean()
    return kk, kk.rolling(d).mean()


def atr(h, l, c, n=14):
    tr = pd.concat([h - l, (h - c.shift()).abs(), (l - c.shift()).abs()], axis=1).max(axis=1)
    return tr.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()


def mansfield(close, bench, n):
    if bench is None:
        return None
    b = bench.reindex(close.index).ffill()
    rs = close / b
    return (rs / rs.rolling(n).mean() - 1) * 100


def indicadores(d, bench, tf, sector=None):
    c, h, l = d["Close"], d["High"], d["Low"]
    ind = {}
    if tf == "D":
        ind["ema50"], ind["ema200"] = ema(c, 50), ema(c, 200)
        ind["vol_ma"] = sma(d["Volume"], 50)
        n_rs = 250
    else:
        ind["ema10"], ind["sma30"] = ema(c, 10), sma(c, 30)
        ind["vol_ma"] = sma(d["Volume"], 10)
        n_rs = 52
    ind["rsi"] = rsi(c)
    ind["macd"], ind["sig"], ind["hist"] = macd(c)
    ind["k"], ind["dd"] = estocastico(h, l, c)
    ind["atr"] = atr(h, l, c)
    ind["rs"] = mansfield(c, bench, n_rs)
    ind["rs_sec"] = mansfield(c, sector, n_rs)
    return ind


# ============================================================================
# Datos: preparacion y semanal
# ============================================================================
def preparar(df):
    """OHLCV con Close ajustado solo por splits (precio real de broker)."""
    return df[["Open", "High", "Low", "Close", "Volume"]].dropna()


def semanal(d):
    w = d.resample("W-FRI").agg({"Open": "first", "High": "max", "Low": "min",
                                  "Close": "last", "Volume": "sum"})
    w = w.dropna(subset=["Close"])
    parcial = bool(len(w) and w.index[-1] > d.index[-1])
    if parcial:
        w = w.iloc[:-1]
    return w, parcial


# ============================================================================
# Etapa de Weinstein (semanal, solo barras cerradas)
# ============================================================================
def etapas_serie(dw, plano=1.0):
    c = dw["Close"]
    m = sma(c, 30)
    pend = (m / m.shift(4) - 1) * 100
    et = pd.Series(np.nan, index=dw.index)
    prev = None
    for i in range(len(dw)):
        if not (np.isfinite(m.iloc[i]) and np.isfinite(pend.iloc[i])):
            continue
        arriba = c.iloc[i] > m.iloc[i]
        p = pend.iloc[i]
        sube, baja = p > plano, p < -plano
        if prev is None:
            e = 2 if (arriba and sube) else 4 if (not arriba and baja) else (3 if arriba else 1)
        elif arriba and sube:
            e = 2
        elif arriba and baja:
            e = 4 if prev == 4 else 3
        elif arriba:
            e = 3 if prev in (2, 3) else 1
        elif sube:
            e = 3 if prev in (2, 3) else 1
        elif baja:
            e = 4
        else:
            e = 3 if prev in (2, 3) else (4 if prev == 4 else 1)
        et.iloc[i] = e
        prev = e
    return et, m, pend


def info_etapa(dw):
    et, m, pend = etapas_serie(dw)
    v = et.dropna()
    if len(v) < 5:
        return {"etapa": "n/d", "semanas": 0, "previa": "", "pend4": np.nan, "dist": np.nan}
    cur = int(v.iloc[-1])
    n = 0
    for x in v.iloc[::-1]:
        if x == cur:
            n += 1
        else:
            break
    previa = int(v.iloc[-n - 1]) if len(v) > n else ""
    return {"etapa": cur, "semanas": n, "previa": previa,
            "pend4": float(pend.iloc[-1]),
            "dist": float((dw["Close"].iloc[-1] / m.iloc[-1] - 1) * 100)}


# ============================================================================
# Pivotes, niveles, estructura, tendencia
# ============================================================================
def pivotes(d, n):
    h, l = d["High"].values, d["Low"].values
    ph, pl = [], []
    for i in range(n, len(d) - n):
        if h[i] == h[i - n:i + n + 1].max():
            ph.append(i)
        if l[i] == l[i - n:i + n + 1].min():
            pl.append(i)
    return ph, pl


def niveles_sr(d, ph, pl, ult, tol=0.015):
    ini = len(d) - ult
    precios = sorted([d["High"].iloc[i] for i in ph if i >= ini] +
                     [d["Low"].iloc[i] for i in pl if i >= ini])
    cl = []
    for p in precios:
        if cl and p <= cl[-1][-1] * (1 + tol):
            cl[-1].append(p)
        else:
            cl.append([p])
    return [(float(np.mean(c)), len(c)) for c in cl if len(c) >= 2]


def soportes_resistencias(niv, c, hi52):
    sop = sorted([(p, t) for p, t in niv if p < c], key=lambda z: -z[0])
    res = sorted([(p, t) for p, t in niv if p > c], key=lambda z: z[0])
    if hi52 > c * 1.001 and not any(abs(p / hi52 - 1) <= 0.015 for p, _ in res):
        res = sorted(res + [(float(hi52), 1)], key=lambda z: z[0])
    return sop, res


def perfil_volumen(d, n=250, bins=40):
    w = d.iloc[-n:]
    tp = (w["High"] + w["Low"] + w["Close"]) / 3
    h, e = np.histogram(tp, bins=bins, weights=w["Volume"])
    mid = (e[:-1] + e[1:]) / 2
    tot = h.sum() or 1.0
    return mid, h / tot


def info_volumen(d, c):
    mid, sh = perfil_volumen(d)
    poc = float(mid[int(np.argmax(sh))])
    hv = [m for m, s in zip(mid, sh) if s >= 1.5 / len(sh)]
    sup = [m for m in hv if m < c * 0.99]
    res = [m for m in hv if m > c * 1.01]
    return poc, (float(max(sup)) if sup else np.nan), (float(min(res)) if res else np.nan)


def estructura(d, ph, pl):
    hs = [d["High"].iloc[i] for i in ph[-2:]]
    ls = [d["Low"].iloc[i] for i in pl[-2:]]
    if len(hs) < 2 or len(ls) < 2:
        return "indefinida", ""
    a = "HH" if hs[1] > hs[0] else "LH"
    b = "HL" if ls[1] > ls[0] else "LL"
    t = "alcista" if (a, b) == ("HH", "HL") else "bajista" if (a, b) == ("LH", "LL") else "mixta"
    return t, f"{a} ({hs[0]:.2f}->{hs[1]:.2f}), {b} ({ls[0]:.2f}->{ls[1]:.2f})"


def tendencia_d(c, e50, e200):
    if not (np.isfinite(e50.iloc[-1]) and np.isfinite(e200.iloc[-1])):
        return "lateral", np.nan
    pend = float(e50.iloc[-1] / e50.iloc[-11] - 1) * 100
    if c > e50.iloc[-1] and e50.iloc[-1] > e200.iloc[-1] and pend > 0:
        return "alcista", pend
    if c < e50.iloc[-1] and e50.iloc[-1] < e200.iloc[-1] and pend < 0:
        return "bajista", pend
    return "lateral", pend


def info_macd(m, sig, hist):
    diff = (m - sig).dropna()
    out = {"hist": float(hist.iloc[-1]), "hist_prev": float(hist.iloc[-2]),
           "cruce": "", "hace": 999, "fecha": None}
    sg = np.sign(diff.values)
    for j in range(len(sg) - 1, 0, -1):
        if sg[j] != sg[j - 1] and sg[j] != 0:
            out.update(cruce="alcista" if sg[j] > 0 else "bajista",
                       hace=len(sg) - 1 - j, fecha=diff.index[j])
            break
    return out


def divergencias(d, rsi_s, ph, pl, ventana, maxedad):
    n = len(d)
    out = []
    ok = np.isfinite
    a = [i for i in ph if i >= n - ventana]
    if len(a) >= 2:
        i1, i2 = a[-2], a[-1]
        r1, r2 = rsi_s.iloc[i1], rsi_s.iloc[i2]
        if (n - 1 - i2 <= maxedad and ok(r1) and ok(r2) and r1 > 55
                and d["High"].iloc[i2] > d["High"].iloc[i1] and r2 < r1 - 2):
            out.append({"tipo": "bajista", "i1": i1, "i2": i2,
                        "p1": d["High"].iloc[i1], "p2": d["High"].iloc[i2],
                        "r1": r1, "r2": r2})
    b = [i for i in pl if i >= n - ventana]
    if len(b) >= 2:
        i1, i2 = b[-2], b[-1]
        r1, r2 = rsi_s.iloc[i1], rsi_s.iloc[i2]
        if (n - 1 - i2 <= maxedad and ok(r1) and ok(r2) and r1 < 45
                and d["Low"].iloc[i2] < d["Low"].iloc[i1] and r2 > r1 + 2):
            out.append({"tipo": "alcista", "i1": i1, "i2": i2,
                        "p1": d["Low"].iloc[i1], "p2": d["Low"].iloc[i2],
                        "r1": r1, "r2": r2})
    return out


def lineas_tendencia(d, ph, pl, ventana, min_sep):
    """Linea de soporte (pivotes bajos crecientes) / resistencia (altos decrecientes), >=3 toques."""
    n = len(d)
    out = []
    cl = d["Close"].values
    for tipo, piv, col in (("soporte", pl, "Low"), ("resistencia", ph, "High")):
        pts = [i for i in piv if i >= n - ventana][-5:]
        if len(pts) < 3:
            continue
        y = d[col].values
        mejor = None
        for a in range(len(pts)):
            for b in range(a + 1, len(pts)):
                i1, i2 = pts[a], pts[b]
                if i2 - i1 < min_sep:
                    continue
                m = (y[i2] - y[i1]) / (i2 - i1)
                if (tipo == "soporte" and m <= 0) or (tipo == "resistencia" and m >= 0):
                    continue
                ks = np.arange(i1, n)
                lv = y[i1] + m * (ks - i1)
                if (lv <= 0).any():
                    continue
                if tipo == "soporte" and (cl[i1:] < lv * 0.99).any():
                    continue
                if tipo == "resistencia" and (cl[i1:] > lv * 1.01).any():
                    continue
                toques = sum(1 for k in pts
                             if abs(y[k] / (y[i1] + m * (k - i1)) - 1) <= 0.012)
                if toques >= 3 and (mejor is None or toques > mejor[0]
                                    or (toques == mejor[0] and i1 < mejor[1])):
                    mejor = (toques, i1, i2, m, float(y[i1]))
        if mejor:
            t, i1, i2, m, y1 = mejor
            out.append({"tipo": tipo, "i1": i1, "p1": y1, "m": m, "toques": t,
                        "valor_hoy": y1 + m * (n - 1 - i1)})
    return out


# ============================================================================
# Patrones de velas (con contexto)
# ============================================================================
def patrones_velas(d):
    o, h, l, c = d["Open"], d["High"], d["Low"], d["Close"]
    body = (c - o).abs()
    rng = (h - l).replace(0, np.nan)
    up = h - np.maximum(o, c)
    lo = np.minimum(o, c) - l
    avgb = body.rolling(14).mean()
    bull, bear = c > o, c < o
    tdown = c.shift(1) < c.shift(6)
    tup = c.shift(1) > c.shift(6)

    def s(x, k):
        return x.shift(k)

    P = {}
    P["Doji"] = ((body <= 0.1 * rng), "neutral")
    P["Martillo"] = ((lo >= 2 * body) & (up <= 0.15 * rng) & tdown, "alcista")
    P["Martillo invertido"] = ((up >= 2 * body) & (lo <= 0.15 * rng) & tdown & bull, "alcista")
    P["Hombre colgado"] = ((lo >= 2 * body) & (up <= 0.15 * rng) & tup, "bajista")
    P["Estrella fugaz"] = ((up >= 2 * body) & (lo <= 0.15 * rng) & tup, "bajista")
    P["Envolvente alcista"] = (s(bear, 1) & bull & (o <= s(c, 1)) & (c >= s(o, 1))
                               & (body > s(body, 1)) & tdown, "alcista")
    P["Envolvente bajista"] = (s(bull, 1) & bear & (o >= s(c, 1)) & (c <= s(o, 1))
                               & (body > s(body, 1)) & tup, "bajista")
    P["Perforante"] = (s(bear, 1) & bull & (o < s(c, 1)) & (c > (s(o, 1) + s(c, 1)) / 2)
                       & (c < s(o, 1)) & tdown, "alcista")
    P["Nube oscura"] = (s(bull, 1) & bear & (o > s(c, 1)) & (c < (s(o, 1) + s(c, 1)) / 2)
                        & (c > s(o, 1)) & tup, "bajista")
    P["Estrella de la manana"] = (s(bear, 2) & (s(body, 2) > s(avgb, 2))
                                  & (s(body, 1) <= 0.3 * s(body, 2)) & bull
                                  & (c > (s(o, 2) + s(c, 2)) / 2), "alcista")
    P["Estrella de la tarde"] = (s(bull, 2) & (s(body, 2) > s(avgb, 2))
                                 & (s(body, 1) <= 0.3 * s(body, 2)) & bear
                                 & (c < (s(o, 2) + s(c, 2)) / 2), "bajista")
    P["Tres soldados blancos"] = (bull & s(bull, 1) & s(bull, 2) & (c > s(c, 1))
                                  & (s(c, 1) > s(c, 2)) & (o > s(o, 1)) & (o < s(c, 1))
                                  & (s(o, 1) > s(o, 2)) & (s(o, 1) < s(c, 2))
                                  & (body > 0.5 * avgb), "alcista")
    P["Tres cuervos negros"] = (bear & s(bear, 1) & s(bear, 2) & (c < s(c, 1))
                                & (s(c, 1) < s(c, 2)) & (o < s(o, 1)) & (o > s(c, 1))
                                & (s(o, 1) < s(o, 2)) & (s(o, 1) > s(c, 2))
                                & (body > 0.5 * avgb), "bajista")
    P["Marubozu alcista"] = (bull & (body >= 0.9 * rng) & (body > 1.5 * avgb), "alcista")
    P["Marubozu bajista"] = (bear & (body >= 0.9 * rng) & (body > 1.5 * avgb), "bajista")
    P["Inside bar"] = ((h <= s(h, 1)) & (l >= s(l, 1)), "neutral")
    return {k: (v[0].fillna(False), v[1]) for k, v in P.items()}


def velas_recientes(t, tf, d, ultimas, niveles, dinam, atr_s):
    rows = []
    P = patrones_velas(d)
    for i in range(max(0, len(d) - ultimas), len(d)):
        a = atr_s.iloc[i]
        for nombre, (ser, sentido) in P.items():
            if not bool(ser.iloc[i]):
                continue
            ref = (d["Low"].iloc[i] if sentido == "alcista"
                   else d["High"].iloc[i] if sentido == "bajista" else d["Close"].iloc[i])
            cand = list(niveles) + [s.iloc[i] for s in dinam if np.isfinite(s.iloc[i])]
            tol = max(0.5 * a, 0.008 * ref) if np.isfinite(a) else 0.01 * ref
            en = any(abs(ref - x) <= tol for x in cand)
            rows.append({"Ticker": t, "TF": tf, "Fecha": d.index[i].date(), "Patron": nombre,
                         "Sentido": sentido, "En nivel": "si" if en else "no", "pos": i})
    return rows


# ============================================================================
# Figuras (con coordenadas para dibujarlas)
# ============================================================================
def figuras(d, tf, ph, pl):
    out = []
    daily = tf == "D"
    n = len(d)
    L = 40 if daily else 20
    prof_max = 0.12 if daily else 0.18
    sep_min = 10 if daily else 6
    edad_max = 40 if daily else 20
    n52 = 252 if daily else 52
    nvol = 50 if daily else 10
    c = d["Close"].iloc[-1]

    # Base / rango lateral y ruptura (la base es el rango previo a la ultima barra)
    if n > L + 2:
        wb = d.iloc[-L - 1:-1]
        top, bot = wb["High"].max(), wb["Low"].min()
        prof = (top - bot) / top
        mv = d["Volume"].iloc[-nvol - 1:-1].mean()
        vr = d["Volume"].iloc[-1] / mv if mv else np.nan
        base = {"tipo": "rect", "x0": n - L - 1, "x1": n - 1, "y0": bot, "y1": top, "nivel": top}
        if c > top and prof <= prof_max + 0.04:
            out.append({**base, "nombre": "Ruptura de base/rango", "vol_ratio": vr,
                        "detalle": f"cierre {c:.2f} > techo {top:.2f} ({L} barras, prof. {prof*100:.1f}%); vol x{vr:.2f}"})
        elif prof <= prof_max and c >= bot * 0.97 and c <= top:
            out.append({**base, "nombre": "Base / rango lateral",
                        "detalle": f"profundidad {prof*100:.1f}% en {L} barras; techo {top:.2f}, suelo {bot:.2f}"})

    # Cerca de maximo 52 semanas
    if n >= n52:
        hi = d["High"].iloc[-n52:].max()
        dist = (c / hi - 1) * 100
        if dist >= -3:
            out.append({"nombre": "Cerca de maximo 52s", "tipo": "nota",
                        "detalle": f"maximo {hi:.2f}, distancia {dist:.1f}%"})

    # Doble suelo / techo
    if len(pl) >= 2:
        i1, i2 = pl[-2], pl[-1]
        l1, l2 = d["Low"].iloc[i1], d["Low"].iloc[i2]
        if abs(l1 - l2) / l1 <= 0.03 and i2 - i1 >= sep_min and n - 1 - i2 <= edad_max:
            cuello = d["High"].iloc[i1:i2 + 1].max()
            if cuello >= max(l1, l2) * 1.08:
                est = "confirmado" if c > cuello else "pendiente de cierre > cuello"
                out.append({"nombre": "Doble suelo", "tipo": "doble", "pts": [(i1, l1), (i2, l2)],
                            "nivel": cuello, "x0": i1,
                            "detalle": f"minimos {l1:.2f}/{l2:.2f}, cuello {cuello:.2f} ({est})"})
    if len(ph) >= 2:
        i1, i2 = ph[-2], ph[-1]
        h1, h2 = d["High"].iloc[i1], d["High"].iloc[i2]
        if abs(h1 - h2) / h1 <= 0.03 and i2 - i1 >= sep_min and n - 1 - i2 <= edad_max:
            cuello = d["Low"].iloc[i1:i2 + 1].min()
            if cuello <= min(h1, h2) * 0.92:
                est = "confirmado" if c < cuello else "pendiente de cierre < cuello"
                out.append({"nombre": "Doble techo", "tipo": "doble", "pts": [(i1, h1), (i2, h2)],
                            "nivel": cuello, "x0": i1,
                            "detalle": f"maximos {h1:.2f}/{h2:.2f}, cuello {cuello:.2f} ({est})"})

    # Triangulos
    ventana_t = 120 if daily else 60
    if (len(ph) >= 3 and len(pl) >= 3 and ph[-3] >= n - ventana_t and pl[-3] >= n - ventana_t
            and n - 1 - max(ph[-1], pl[-1]) <= edad_max):
        iH, iL = ph[-3:], pl[-3:]
        H = [d["High"].iloc[i] for i in iH]
        Lw = [d["Low"].iloc[i] for i in iL]
        flat_h = (max(H) - min(H)) / np.mean(H) <= 0.02
        flat_l = (max(Lw) - min(Lw)) / np.mean(Lw) <= 0.02
        lineas = [((iH[0], H[0]), (iH[2], H[2])), ((iL[0], Lw[0]), (iL[2], Lw[2]))]
        nom = det = None
        if flat_h and Lw[0] < Lw[1] < Lw[2]:
            nom, det = "Triangulo ascendente", f"resistencia ~{np.mean(H):.2f}, minimos crecientes"
        elif flat_l and H[0] > H[1] > H[2]:
            nom, det = "Triangulo descendente", f"soporte ~{np.mean(Lw):.2f}, maximos decrecientes"
        elif H[0] > H[1] > H[2] and Lw[0] < Lw[1] < Lw[2]:
            nom, det = "Triangulo simetrico / cuna", "maximos decrecientes y minimos crecientes"
        if nom:
            out.append({"nombre": nom, "tipo": "lineas", "lineas": lineas, "detalle": det})

    # Bandera alcista (solo diario)
    if daily and n > 45:
        seg = d.iloc[-40:]
        hi_i = int(seg["High"].values.argmax())
        if hi_i >= 5:
            pre, cons = seg.iloc[:hi_i + 1], seg.iloc[hi_i:]
            lo_i = int(pre["Low"].values.argmin())
            base = pre["Low"].min()
            hi = seg["High"].max()
            gain = hi / base - 1
            retr = (hi - cons["Low"].min()) / (hi - base) if hi > base else 1
            rango = (cons["High"].max() - cons["Low"].min()) / hi
            if (gain >= 0.15 and hi_i - lo_i <= 20 and 4 <= len(cons) - 1 <= 20
                    and retr <= 0.5 and rango <= 0.12):
                cont = cons["Volume"].mean() < pre["Volume"].iloc[-10:].mean()
                off = n - 40
                out.append({"nombre": "Bandera alcista (posible)", "tipo": "bandera",
                            "mastil": ((off + lo_i, base), (off + hi_i, hi)),
                            "x0": off + hi_i, "x1": n - 1,
                            "y0": cons["Low"].min(), "y1": cons["High"].max(), "nivel": hi,
                            "detalle": f"mastil +{gain*100:.0f}%, retroceso {retr*100:.0f}%, "
                                       f"volumen {'decreciente' if cont else 'no decreciente'}"})

    # Taza (solo semanal)
    if not daily and n >= 40:
        N = min(60, n)
        w = d.iloc[-N:]
        left_i = int(w["High"].values[:N // 3].argmax())
        left = w["High"].iloc[left_i]
        bot_i = int(w["Low"].values[left_i:].argmin()) + left_i
        bottom = w["Low"].iloc[bot_i]
        prof = (left - bottom) / left
        right_i = bot_i + int(w["High"].values[bot_i:].argmax())
        right = w["High"].iloc[right_i]
        if (0.15 <= prof <= 0.35 and bot_i - left_i >= 10 and N - 1 - bot_i >= 6
                and right >= 0.92 * left and 0.92 * left <= c <= 1.05 * left):
            off = n - N
            out.append({"nombre": "Taza (posible)", "tipo": "taza",
                        "pts": [(off + left_i, left), (off + bot_i, bottom), (off + right_i, right)],
                        "nivel": left,
                        "detalle": f"borde izq {left:.2f}, fondo {bottom:.2f} ({prof*100:.0f}%), pivote ~{left:.2f}"})
    return out


# ============================================================================
# Plan de compra
# ============================================================================
def objetivos_figuras(figs, tf):
    """Objetivo por altura proyectada de la figura (movimiento medido)."""
    out = []
    for f in figs:
        n, t = f["nombre"], f["tipo"]
        if t == "rect":
            out.append((f["y1"] + (f["y1"] - f["y0"]), f"{n} {tf}"))
        elif t == "bandera":
            (_, pa), (_, pb) = f["mastil"]
            out.append((f["nivel"] + (pb - pa), f"{n} {tf}"))
        elif t == "taza":
            left, bottom = f["pts"][0][1], f["pts"][1][1]
            out.append((left + (left - bottom), f"{n} {tf}"))
        elif t == "doble" and n == "Doble suelo":
            lo = min(p for _, p in f["pts"])
            out.append((f["nivel"] + (f["nivel"] - lo), f"{n} {tf}"))
    return out


def plan_compra(c, a, sop, res, lows, cfg, objetivos=(), trail=np.nan):
    """Entrada, stop, TP1/TP2 y acciones. Devuelve (plan | None, motivo)."""
    if not np.isfinite(a) or a <= 0:
        return None, "ATR no disponible"
    r1 = res[0][0] if res else None
    if r1 is not None and 0 < (r1 / c - 1) <= 0.05:
        modo = "ruptura sobre %.2f (orden stop de compra)" % r1
        e_lo, e_hi = r1 + 0.05 * a, r1 + 0.30 * a
    else:
        modo = "a mercado (cierre actual)"
        e_lo, e_hi = c - 0.25 * a, c
    entrada = e_hi
    cand = [x for x in [p for p, _ in sop] + list(lows) if x <= e_lo - 0.5 * a]
    stop_tipo = "bajo soporte/pivote"
    if cand:
        stop = max(cand) - 0.25 * a
    else:
        stop, stop_tipo = e_lo - 2 * a, "2xATR"
    if (entrada - stop) / entrada > 0.12:
        stop, stop_tipo = e_lo - 2 * a, "2xATR (soporte lejano)"
    riesgo_acc = entrada - stop
    if riesgo_acc <= 0 or riesgo_acc / entrada > 0.12:
        return None, "stop a mas del 12% de la entrada"
    niveles = [p for p, _ in res if p > entrada * 1.005]
    if niveles:
        tp1, tp1_tipo = niveles[0], "resistencia"
    else:
        tp1, tp1_tipo = entrada + 1.5 * riesgo_acc, "multiplo 1.5R (sin resistencia por encima)"
    cand = [(v, nm) for v, nm in sorted(objetivos) if v > tp1 * 1.005]
    if cand:
        tp2, tp2_tipo = cand[0][0], f"altura de figura ({cand[0][1]})"
    elif len(niveles) >= 2 and niveles[1] > tp1 * 1.005:
        tp2, tp2_tipo = niveles[1], "resistencia"
    else:
        tp2, tp2_tipo = max(entrada + 3 * riesgo_acc, tp1 + riesgo_acc), "multiplo 3R (sin figura ni resistencia)"
    tp_tipo = f"TP1 {tp1_tipo}; TP2 {tp2_tipo}"
    n_r = int(cfg["riesgo_max"] // riesgo_acc)
    n_c = int(cfg["max_pos"] // entrada)
    n = min(n_r, n_c)
    if n < 1:
        return None, "no cabe ni 1 accion"
    return {"modo": modo, "e_lo": e_lo, "e_hi": e_hi, "entrada": entrada, "stop": stop,
            "stop_tipo": stop_tipo, "riesgo_acc": riesgo_acc, "tp1": tp1, "tp2": tp2,
            "tp_tipo": tp_tipo, "tp2_tipo": tp2_tipo, "trail": trail, "rr1": (tp1 - entrada) / riesgo_acc,
            "rr2": (tp2 - entrada) / riesgo_acc, "n_riesgo": n_r, "n_capital": n_c,
            "acciones": n, "limite": "riesgo" if n_r <= n_c else "capital",
            "capital": n * entrada, "riesgo_real": n * riesgo_acc}, ""


# ============================================================================
# Analisis por ticker
# ============================================================================
def analizar_tf(t, tf, d, ind, cfg):
    n_piv = 5 if tf == "D" else 3
    ph, pl = pivotes(d, n_piv)
    niv = niveles_sr(d, ph, pl, 250 if tf == "D" else 150)
    c = d["Close"].iloc[-1]
    hi52 = d["High"].iloc[-(252 if tf == "D" else 52):].max()
    lo52 = d["Low"].iloc[-(252 if tf == "D" else 52):].min()
    sop, res = soportes_resistencias(niv, c, hi52)
    dinam = ([ind["ema50"], ind["ema200"]] if tf == "D" else [ind["ema10"], ind["sma30"]])
    ult = cfg["marcar"] if tf == "D" else 6
    vel = velas_recientes(t, tf, d, ult, [p for p, _ in niv], dinam, ind["atr"])
    tl = lineas_tendencia(d, ph, pl, 120 if tf == "D" else 78, 8 if tf == "D" else 4)
    dv = divergencias(d, ind["rsi"], ph, pl, 60 if tf == "D" else 30, 25 if tf == "D" else 12)
    fgs = figuras(d, tf, ph, pl)
    return {"d": d, "ind": ind, "ph": ph, "pl": pl, "niv": niv, "sop": sop, "res": res,
            "vel": vel, "tl": tl, "div": dv, "fig": fgs, "hi52": hi52, "lo52": lo52}


def val_rs(s, k):
    if s is None:
        return np.nan
    v = s.dropna()
    return float(v.iloc[-1 - k]) if len(v) > k else np.nan


def cruce_cero(s, ventana=13):
    if s is None:
        return None
    v = s.dropna().iloc[-ventana - 1:]
    sg = np.sign(v.values)
    for j in range(len(sg) - 1, 0, -1):
        if sg[j] != sg[j - 1] and sg[j] != 0:
            return len(sg) - 1 - j
    return None


def analizar(t, raw, bench_d, bench_w, sec_d, sec_w, sector_nombre, cfg):
    dd = preparar(raw)
    if len(dd) < MIN_BARRAS_D:
        return None, f"datos insuficientes ({len(dd)} barras diarias, minimo {MIN_BARRAS_D})"
    dw, parcial = semanal(dd)
    if len(dw) < MIN_BARRAS_W:
        return None, f"datos insuficientes ({len(dw)} barras semanales)"
    avisos = []
    if len(dd) < 500:
        avisos.append(f"{t}: historia corta ({len(dd)} barras diarias)")
    gaps = dd.index.to_series().diff().dt.days.iloc[-250:]
    if (gaps > 7).any():
        avisos.append(f"{t}: hueco de mas de 7 dias en el ultimo año")
    if (dd["High"].iloc[-250:] < dd["Low"].iloc[-250:]).any():
        avisos.append(f"{t}: barras con High < Low")

    ind_d = indicadores(dd, bench_d, "D", sec_d)
    ind_w = indicadores(dw, bench_w, "W", sec_w)
    D = analizar_tf(t, "D", dd, ind_d, cfg)
    W = analizar_tf(t, "W", dw, ind_w, cfg)
    et = info_etapa(dw)
    c = float(dd["Close"].iloc[-1])
    a = float(ind_d["atr"].iloc[-1])
    tend, pend50 = tendencia_d(c, ind_d["ema50"], ind_d["ema200"])
    estr, estr_txt = estructura(dd, D["ph"], D["pl"])
    mac = info_macd(ind_d["macd"], ind_d["sig"], ind_d["hist"])
    poc, hvn_s, hvn_r = info_volumen(dd, c)
    mv = dd["Volume"].iloc[-51:-1].mean()
    lows_rec = [dd["Low"].iloc[i] for i in D["pl"] if i >= len(dd) - 60 and dd["Low"].iloc[i] < c]
    rs = ind_w["rs"]

    r = {"ticker": t, "dd": dd, "dw": dw, "parcial": parcial, "avisos": avisos, "tf": {"D": D, "W": W},
         "fecha": dd.index[-1], "fecha_w": dw.index[-1], "cierre": c, "atr": a,
         "atr_pct": a / c * 100, "sector": sector_nombre,
         "etapa": et["etapa"], "etapa_sem": et["semanas"], "etapa_prev": et["previa"],
         "pend_sma30": et["pend4"], "dist_sma30": et["dist"],
         "tend_d": tend, "pend_ema50": pend50, "estructura": estr, "estructura_txt": estr_txt,
         "ema50": float(ind_d["ema50"].iloc[-1]), "ema200": float(ind_d["ema200"].iloc[-1]),
         "rsi_d": float(ind_d["rsi"].iloc[-1]), "rsi_w": float(ind_w["rsi"].iloc[-1]),
         "macd": mac, "div_d": D["div"], "div_w": W["div"],
         "rs_sp": val_rs(rs, 0), "rs_sp_4": val_rs(rs, 4), "rs_sp_13": val_rs(rs, 13),
         "rs_cruce": cruce_cero(rs), "rs_sec": val_rs(ind_w["rs_sec"], 0),
         "dist_max52": (c / D["hi52"] - 1) * 100,
         "vol_ratio": float(dd["Volume"].iloc[-1] / mv) if mv else np.nan,
         "poc": poc, "hvn_sop": hvn_s, "hvn_res": hvn_r}
    objetivos = objetivos_figuras(D["fig"], "D") + objetivos_figuras(W["fig"], "S")
    sma30w = float(ind_w["sma30"].iloc[-1])
    plan, motivo = plan_compra(c, a, D["sop"], D["res"], lows_rec, cfg, objetivos, sma30w)
    if np.isfinite(r["vol_ratio"]) and r["vol_ratio"] < 0.35:
        avisos.append(f"{t}: volumen de la ultima barra x{r['vol_ratio']:.2f} de la media 50d "
                      f"(sesion corta o dato parcial)")
    r["plan"], r["plan_motivo"] = plan, motivo
    r["score"] = puntuar(r)
    clasificar(r)
    return r, avisos


def puntuar(r):
    s = 0
    s += {2: 25, 3: 5, 1: 8, 4: 0}.get(r["etapa"], 0)
    s += {"alcista": 15, "lateral": 6, "bajista": 0}[r["tend_d"]]
    rs = r["rs_sp"]
    if np.isfinite(rs):
        s += 15 if rs > 0 else 0
        if np.isfinite(r["rs_sp_4"]) and rs > r["rs_sp_4"]:
            s += 5
    m = r["macd"]
    if m["hist"] > 0:
        s += 5
    if m["cruce"] == "alcista" and m["hace"] <= 10:
        s += 5
    if m["cruce"] == "bajista" and m["hace"] <= 5:
        s -= 5
    if 50 <= r["rsi_d"] <= 70:
        s += 5
    elif r["rsi_d"] > 75:
        s -= 5
    for dv in r["div_d"]:
        if dv["tipo"] == "bajista":
            s -= 10
        elif r["tend_d"] != "bajista":
            s += 5
    # Ruptura de base y volumen NO puntuan: en el backtest (S&P 500, 9 años) no mejoraron a
    # etapa 2 + RS>0 y su variante completa fue la peor (ver backtest.py).
    if r["dist_max52"] >= -10:
        s += 5
    if np.isfinite(r["dist_sma30"]) and r["dist_sma30"] > 25:
        s -= 10
    if r["plan"]:
        rr2 = r["plan"]["rr2"]
        s += 15 if rr2 >= 3 else 10 if rr2 >= 2 else 5 if rr2 >= 1.5 else 0
    n = len(r["dd"])
    if any(v["Sentido"] == "alcista" and v["En nivel"] == "si" and v["pos"] >= n - 3
           for v in r["tf"]["D"]["vel"]):
        s += 5
    if any(v["Sentido"] == "bajista" and v["En nivel"] == "si" and v["pos"] >= n - 3
           for v in r["tf"]["D"]["vel"]):
        s -= 5
    return int(max(0, min(100, s)))


def clasificar(r):
    et, tend, rs = r["etapa"], r["tend_d"], r["rs_sp"]
    razones = []
    if et != 2:
        razones.append(f"etapa semanal {et}")
    if tend != "alcista":
        razones.append(f"tendencia diaria {tend}")
    if not (np.isfinite(rs) and rs > 0):
        razones.append("RS Mansfield no positivo")
    if np.isfinite(r["dist_sma30"]) and r["dist_sma30"] > 25:
        razones.append(f"extendida (+{r['dist_sma30']:.0f}% sobre SMA30s)")
    if any(v["tipo"] == "bajista" for v in r["div_d"]):
        razones.append("divergencia bajista RSI")
    p = r["plan"]
    if p is None:
        razones.append("sin plan valido: " + r["plan_motivo"])
    elif p["rr1"] < 1:
        razones.append(f"resistencia cercana: R:R TP1 {p['rr1']:.1f} < 1")
    elif p["rr2"] < 2:
        razones.append(f"R:R TP2 {p['rr2']:.1f} < 2")
    if et == 4 and tend == "bajista" and not (np.isfinite(rs) and rs > 0):
        r["dir"] = "Venta"
        r["razon"] = "etapa 4, tendencia diaria bajista y RS negativo"
    elif not razones:
        r["dir"] = "Compra"
        r["razon"] = "etapa 2, tendencia alcista, RS>0, R:R>=2"
    else:
        r["dir"] = "Esperar"
        r["razon"] = "; ".join(razones)
    m, n = r["macd"], len(r["dd"])
    al = []
    if m["cruce"] == "bajista" and m["hace"] <= 5:
        al.append(f"MACD cruce bajista hace {m['hace']} barras")
    if any(v["Sentido"] == "bajista" and v["pos"] >= n - 3 and v["Patron"] != "Inside bar"
           for v in r["tf"]["D"]["vel"]):
        al.append("vela bajista reciente")
    if np.isfinite(r["rsi_d"]) and r["rsi_d"] > 75:
        al.append("RSI sobrecomprado")
    r["alertas"] = "; ".join(al)
    res = r["tf"]["D"]["res"]
    r["activar"] = res[0][0] if res else np.nan
    if r["dir"] != "Compra":
        r["plan_usable"] = None
    else:
        r["plan_usable"] = r["plan"]


# ============================================================================
# Graficos
# ============================================================================
ESTILO_MAS = {
    "ema50": ("EMA50", "#ff7f0e", 1.2, "-"), "ema200": ("EMA200", "#7b3fa0", 1.3, "--"),
    "ema10": ("EMA10", "#1f77b4", 1.1, "-"), "sma30": ("SMA30 (Weinstein)", "#2ca02c", 1.5, "--"),
}
PRIORIDAD_FIG = ["Ruptura de base/rango", "Bandera alcista (posible)", "Taza (posible)",
                 "Triangulo ascendente", "Triangulo descendente", "Triangulo simetrico / cuna",
                 "Doble suelo", "Doble techo", "Base / rango lateral"]


def dibujar_figura(ax, f, off, L, color="#1f4e79"):
    def X(i):
        return i - off

    t = f["tipo"]
    if t in ("rect", "bandera"):
        x0, x1 = max(X(f["x0"]), 0), X(f["x1"])
        ax.add_patch(plt.Rectangle((x0, f["y0"]), x1 - x0, f["y1"] - f["y0"], fc=color,
                                   alpha=0.09, ec=color, lw=0.9, ls="--", zorder=1))
        if t == "bandera":
            (a, pa), (b, pb) = f["mastil"]
            ax.plot([X(a), X(b)], [pa, pb], color=color, lw=1.6, zorder=3)
        ax.text(x0 + 0.5, f["y1"], f["nombre"], fontsize=8, color=color, va="bottom", fontweight="bold")
    elif t == "doble":
        xs = [X(i) for i, _ in f["pts"]]
        ys = [p for _, p in f["pts"]]
        ax.plot(xs, ys, "o", color=color, ms=6, zorder=4)
        ax.hlines(f["nivel"], max(X(f["x0"]), 0), L - off + 2, colors=color, ls="--", lw=1.0)
        ax.text(max(X(f["x0"]), 0), f["nivel"], " " + f["nombre"] + " (cuello)", fontsize=8,
                color=color, va="bottom", fontweight="bold")
    elif t == "lineas":
        for (i1, p1), (i2, p2) in f["lineas"]:
            m = (p2 - p1) / (i2 - i1)
            xs = np.arange(max(i1, off), L + 3)
            ax.plot(xs - off, p1 + m * (xs - i1), color=color, lw=1.2, ls="-", zorder=3)
        (i1, p1), _ = f["lineas"][0]
        ax.text(max(X(i1), 0), p1, " " + f["nombre"], fontsize=8, color=color, va="bottom", fontweight="bold")
    elif t == "taza":
        xs = [X(i) for i, _ in f["pts"]]
        ys = [p for _, p in f["pts"]]
        ax.plot(xs, ys, "o-", color=color, lw=1.3, ms=5, zorder=3)
        ax.hlines(f["nivel"], max(xs[0], 0), L - off + 2, colors=color, ls="--", lw=1.0)
        ax.text(max(xs[0], 0), f["nivel"], " " + f["nombre"] + " (pivote)", fontsize=8,
                color=color, va="bottom", fontweight="bold")


def dibujar(pdf, r, tf, cfg):
    T = r["tf"][tf]
    d, ind = T["d"], T["ind"]
    n = min(cfg["dias"] if tf == "D" else cfg["semanas"], len(d))
    off = len(d) - n
    w = d.iloc[-n:]
    idx = w.index
    x = np.arange(n)
    o, h, l, c = (w[k].values for k in ("Open", "High", "Low", "Close"))
    det = {"limpio": 0, "normal": 1, "completo": 2}[cfg["detalle"]]
    k_sr = (1, 2, 4)[det]
    k_tl = (1, 2, 2)[det]
    k_fig = (1, 1, 9)[det]

    paneles = ["precio", "vol", "rsi", "macd"]
    if ind.get("rs") is not None:
        paneles.append("rs")
    if cfg["estoc"]:
        paneles.append("estoc")
    ratios = {"precio": 7, "vol": 1.6, "rsi": 1.9, "macd": 1.9, "rs": 1.7, "estoc": 1.7}
    fig, axs = plt.subplots(len(paneles), 1, figsize=(16.5, 11.7), sharex=True,
                            gridspec_kw={"height_ratios": [ratios[p] for p in paneles], "hspace": 0.06})
    axs = dict(zip(paneles, np.atleast_1d(axs)))
    ax = axs["precio"]

    col = np.where(c >= o, VERDE, ROJO)
    ax.vlines(x, l, h, colors=col, linewidth=0.9, zorder=2)
    ax.bar(x, np.abs(c - o) + 1e-9, bottom=np.minimum(o, c), width=0.7, color=col, edgecolor=col, zorder=3)

    for k, (nombre, color, lw, ls) in ESTILO_MAS.items():
        if k in ind:
            s = ind[k].iloc[-n:].values
            if np.isfinite(s).any():
                ax.plot(x, s, color=color, lw=lw, ls=ls, label=nombre, zorder=4)

    ax.axhline(T["hi52"], color="grey", ls="--", lw=0.8)
    ax.axhline(T["lo52"], color="grey", ls="--", lw=0.8)
    ax.text(n - 0.5, T["hi52"], f" Max52 {T['hi52']:.2f}", va="bottom", ha="right", fontsize=8, color="grey")
    ax.text(n - 0.5, T["lo52"], f" Min52 {T['lo52']:.2f}", va="top", ha="right", fontsize=8, color="grey")
    cierre = c[-1]
    for p, tq in T["sop"][:k_sr]:
        ax.axhline(p, color="#2e7d32", ls=":", lw=1.0)
        ax.text(0, p, f" S {p:.2f} ({tq})", fontsize=7.5, va="bottom", color="#2e7d32")
    for p, tq in T["res"][:k_sr]:
        ax.axhline(p, color="#b71c1c", ls=":", lw=1.0)
        ax.text(0, p, f" R {p:.2f} ({tq})", fontsize=7.5, va="bottom", color="#b71c1c")

    L = len(d)
    for tl in T["tl"][:k_tl]:
        xs = np.arange(max(tl["i1"], off), L + 3)
        color = VERDE if tl["tipo"] == "soporte" else ROJO
        ax.plot(xs - off, tl["p1"] + tl["m"] * (xs - tl["i1"]), color=color, lw=1.3, ls="-.", zorder=3)
        xe = L + 2 - off
        ax.text(xe, tl["p1"] + tl["m"] * (L + 2 - tl["i1"]), f" tend. {tl['tipo']} ({tl['toques']}t)",
                fontsize=7, color=color, va="center")

    if det == 2:
        for i in T["ph"]:
            if i >= off:
                ax.plot(i - off, h[i - off], "v", color=ROJO, ms=4, zorder=5)
        for i in T["pl"]:
            if i >= off:
                ax.plot(i - off, l[i - off], "^", color=VERDE, ms=4, zorder=5)

    figs = [f for nom in PRIORIDAD_FIG for f in T["fig"] if f["nombre"] == nom and f["tipo"] != "nota"]
    for f in figs[:k_fig]:
        dibujar_figura(ax, f, off, L)

    if det >= 1:
        for dv in T["div"]:
            if dv["i1"] >= off:
                color = VERDE if dv["tipo"] == "alcista" else ROJO
                ax.plot([dv["i1"] - off, dv["i2"] - off], [dv["p1"], dv["p2"]], color=color, lw=1.6,
                        ls="--", zorder=5)

    rng_y = np.nanmax(h) - np.nanmin(l)
    marcas = []
    for v in T["vel"]:
        p = v["pos"] - off
        if not (0 <= p < n and p >= n - cfg["marcar"]):
            continue
        ruido = v["Patron"] in ("Doji", "Inside bar")
        if det == 0 and (v["En nivel"] != "si" or ruido):
            continue
        if det == 1 and ruido and v["En nivel"] != "si":
            continue
        marcas.append(v)
    marcas = marcas[-(3, 4, 99)[det]:]
    for v in marcas:
        p = v["pos"] - off
        if v["Sentido"] == "alcista":
            ax.annotate(v["Patron"], (p, l[p]), xytext=(p, l[p] - 0.04 * rng_y), fontsize=7, color=VERDE,
                        ha="center", va="top", rotation=90,
                        arrowprops=dict(arrowstyle="-", color=VERDE, lw=0.5))
        else:
            ax.annotate(v["Patron"], (p, h[p]), xytext=(p, h[p] + 0.04 * rng_y), fontsize=7,
                        color=ROJO if v["Sentido"] == "bajista" else "#444444", ha="center", va="bottom",
                        rotation=90,
                        arrowprops=dict(arrowstyle="-", color=ROJO, lw=0.5) if v["Sentido"] == "bajista" else None)

    ymin, ymax = np.nanmin(l), np.nanmax(h)
    plan = r["plan_usable"] if tf == "D" else None
    if plan:
        ymin = min(ymin, plan["stop"])
        if plan["tp1"] <= ymax * 1.08:
            ymax = max(ymax, plan["tp1"])
        ax.axhspan(plan["e_lo"], plan["e_hi"], color="#1f77b4", alpha=0.18, zorder=1)
        ax.hlines(plan["stop"], n - 30, n + 3, colors=ROJO, lw=1.4, zorder=5)
        for k, key in (("TP1", "tp1"), ("TP2", "tp2")):
            if plan[key] <= ymax * 1.02:
                ax.hlines(plan[key], n - 30, n + 3, colors=VERDE, lw=1.4, zorder=5)
        txt = (f"PLAN ({plan['modo']})\nEntrada {plan['e_lo']:.2f}-{plan['e_hi']:.2f}\n"
               f"Stop {plan['stop']:.2f} ({plan['stop_tipo']})\n"
               f"TP1 {plan['tp1']:.2f} (R {plan['rr1']:.1f})  TP2 {plan['tp2']:.2f} (R {plan['rr2']:.1f})\n"
               f"Salida: cierre semanal < SMA30s ({f2(plan['trail'])})\n"
               f"{plan['acciones']} acc · capital {plan['capital']:.0f} · riesgo {plan['riesgo_real']:.0f}")
        ax.text(0.01, 0.90, txt, transform=ax.transAxes, fontsize=8.5, va="top", family="monospace",
                bbox=dict(boxstyle="round,pad=0.4", fc="white", ec="#1f77b4", alpha=0.9), zorder=10)
    ax.set_ylim(ymin - 0.10 * (ymax - ymin), ymax + 0.20 * (ymax - ymin))
    ax.set_xlim(-1, n + 6)
    ax.annotate(f"{cierre:.2f}", (n - 1, cierre), xytext=(n + 0.5, cierre), fontsize=9, fontweight="bold",
                color="white", va="center", bbox=dict(boxstyle="round,pad=0.2", fc="black"))
    ax.legend(loc="upper left", fontsize=8, ncol=5, frameon=False)
    ax.grid(alpha=0.25)
    tfn = "Diario" if tf == "D" else "Semanal"
    ax.set_title(f"{r['ticker']} · {tfn} · ultima barra cerrada {idx[-1]:%d/%m/%Y} · cierre {cierre:.2f} · "
                 f"Etapa semanal {r['etapa']} · {r['dir']} ({r['score']})",
                 fontsize=13, fontweight="bold", loc="left")

    a = axs["vol"]
    a.bar(x, w["Volume"].values, width=0.7, color=col, alpha=0.6)
    a.plot(x, ind["vol_ma"].iloc[-n:].values, color="black", lw=1.0,
           label="Media vol " + ("50d" if tf == "D" else "10s"))
    a.set_ylabel("Volumen", fontsize=8)
    a.legend(loc="upper left", fontsize=7, frameon=False)
    a.grid(alpha=0.25)

    a = axs["rsi"]
    a.plot(x, ind["rsi"].iloc[-n:].values, color="#6a3d9a", lw=1.2)
    for y, ls in ((70, "--"), (50, ":"), (30, "--")):
        a.axhline(y, color="grey", ls=ls, lw=0.8)
    if det >= 1:
        for dv in T["div"]:
            if dv["i1"] >= off:
                color = VERDE if dv["tipo"] == "alcista" else ROJO
                a.plot([dv["i1"] - off, dv["i2"] - off], [dv["r1"], dv["r2"]], color=color, lw=1.6,
                       ls="--", marker="o", ms=4)
    a.set_ylim(10, 90)
    a.set_ylabel("RSI 14", fontsize=8)
    a.grid(alpha=0.25)

    a = axs["macd"]
    hist = ind["hist"].iloc[-n:].values
    a.bar(x, hist, width=0.7, color=np.where(hist >= 0, VERDE, ROJO), alpha=0.6)
    a.plot(x, ind["macd"].iloc[-n:].values, color="#1f77b4", lw=1.1)
    a.plot(x, ind["sig"].iloc[-n:].values, color="#ff7f0e", lw=1.1)
    a.axhline(0, color="grey", lw=0.8)
    a.set_ylabel("MACD 12/26/9", fontsize=8)
    a.grid(alpha=0.25)

    if "rs" in axs:
        a = axs["rs"]
        rr = ind["rs"].iloc[-n:].values
        a.plot(x, rr, color="#00695c", lw=1.3)
        a.fill_between(x, 0, rr, where=np.nan_to_num(rr) >= 0, color=VERDE, alpha=0.15)
        a.fill_between(x, 0, rr, where=np.nan_to_num(rr) < 0, color=ROJO, alpha=0.15)
        a.axhline(0, color="grey", lw=0.8)
        a.set_ylabel("RS Mansfield", fontsize=8)
        a.grid(alpha=0.25)

    if "estoc" in axs:
        a = axs["estoc"]
        a.plot(x, ind["k"].iloc[-n:].values, color="#1f77b4", lw=1.1)
        a.plot(x, ind["dd"].iloc[-n:].values, color="#ff7f0e", lw=1.1)
        for y in (20, 80):
            a.axhline(y, color="grey", ls="--", lw=0.8)
        a.set_ylim(0, 100)
        a.set_ylabel("Estoc 14,3,3", fontsize=8)
        a.grid(alpha=0.25)

    ult = axs[paneles[-1]]
    pos, lab = [], []
    for i in range(n):
        if i == 0 or idx[i].month != idx[i - 1].month:
            if tf == "D" or idx[i].month in (1, 4, 7, 10):
                pos.append(i)
                lab.append(idx[i].strftime("%b %y"))
    if len(pos) > 1 and pos[1] - pos[0] < 8:
        pos, lab = pos[1:], lab[1:]
    ult.set_xticks(pos)
    ult.set_xticklabels(lab, fontsize=8)
    fig.text(0.995, 0.005, "Generado por generador_graficos.py", ha="right", fontsize=6, color="grey")
    fig.subplots_adjust(left=0.05, right=0.97, top=0.95, bottom=0.04)
    pdf.savefig(fig)
    plt.close(fig)


# ============================================================================
# Salidas
# ============================================================================
def f2(x, nd=2):
    try:
        return "n/d" if x is None or not np.isfinite(x) else f"{x:.{nd}f}"
    except Exception:
        return "n/d"


def pct(a, b):
    return (a / b - 1) * 100


def ficha(r, cfg):
    t, c = r["ticker"], r["cierre"]
    D, W = r["tf"]["D"], r["tf"]["W"]
    p, m = r["plan_usable"], r["macd"]
    L = [f"## {t} — {r['dir']} · score {r['score']} · cierre {c:.2f}"]
    L.append(f"- Datos: ultima barra diaria cerrada {r['fecha']:%Y-%m-%d}; ultima semanal cerrada "
             f"{r['fecha_w']:%Y-%m-%d}" + ("; semana en curso descartada" if r["parcial"] else "")
             + (f"; sector {r['sector']}" if r.get("sector") else ""))
    L.append(f"- Etapa Weinstein: {r['etapa']} ({r['etapa_sem']} sem; previa {r['etapa_prev'] or 'n/d'}); "
             f"pendiente SMA30s {f2(r['pend_sma30'])}% en 4 sem; distancia {f2(r['dist_sma30'])}%")
    L.append(f"- Tendencia diaria: {r['tend_d']}; estructura {r['estructura']}: {r['estructura_txt']}")
    L.append(f"- Medias D: EMA50 {r['ema50']:.2f} ({pct(c, r['ema50']):+.1f}%, pendiente {f2(r['pend_ema50'])}%/10b), "
             f"EMA200 {r['ema200']:.2f} ({pct(c, r['ema200']):+.1f}%)")
    dvs = "; ".join(f"{v['tipo']} ({D['d'].index[v['i1']]:%d/%m}->{D['d'].index[v['i2']]:%d/%m})" for v in r["div_d"]) or "ninguna"
    L.append(f"- RSI D {r['rsi_d']:.1f} / S {r['rsi_w']:.1f}; divergencias D: {dvs}")
    sh = "positivo" if m["hist"] > 0 else "negativo"
    sh += " creciente" if m["hist"] > m["hist_prev"] else " decreciente"
    cr = f"ultimo cruce {m['cruce']} hace {m['hace']} barras ({m['fecha']:%Y-%m-%d})" if m["cruce"] else "sin cruce reciente"
    L.append(f"- MACD D: histograma {sh} ({m['hist']:.3f}); {cr}")
    cz = f"cruce de cero hace {r['rs_cruce']} sem" if r["rs_cruce"] is not None else "sin cruce de cero en 13 sem"
    L.append(f"- RS Mansfield S vs S&P: {f2(r['rs_sp'])} (hace 4s {f2(r['rs_sp_4'])}, 13s {f2(r['rs_sp_13'])}); {cz}; "
             f"vs sector {f2(r['rs_sec'])}; percentil RS {f2(r.get('rs_pct'), 0)}")
    for tf, T in (("D", D), ("W", W)):
        sp = ", ".join(f"{p_:.2f} ({pct(p_, c):+.1f}%, {tq}t)" for p_, tq in T["sop"][:2]) or "ninguno"
        rs_ = ", ".join(f"{p_:.2f} ({pct(p_, c):+.1f}%, {tq}t)" for p_, tq in T["res"][:2]) or "ninguna (maximos)"
        L.append(f"- Niveles {tf}: soportes {sp}; resistencias {rs_}; Max52 {T['hi52']:.2f} ({pct(c, T['hi52']):+.1f}%), Min52 {T['lo52']:.2f}")
    L.append(f"- Volumen (D, 250 barras): POC {r['poc']:.2f}; nodo alto bajo precio {f2(r['hvn_sop'])}; sobre precio {f2(r['hvn_res'])}; "
             f"volumen ultima barra x{f2(r['vol_ratio'])} de la media 50d")
    tls = "; ".join(f"{tl['tipo']} {tl['toques']} toques, valor hoy {tl['valor_hoy']:.2f}" for tf in ("D", "W") for tl in r["tf"][tf]["tl"]) or "ninguna con 3 toques"
    L.append(f"- Lineas de tendencia: {tls}")
    n = len(D["d"])
    vs = [f"{v['Fecha']} {v['Patron']} ({v['Sentido']}, {'en nivel' if v['En nivel'] == 'si' else 'sin nivel'})"
          for v in D["vel"] if v["pos"] >= n - 5]
    L.append("- Velas D (5 barras): " + ("; ".join(vs) or "ninguna"))
    nw = len(W["d"])
    vw = [f"{v['Fecha']} {v['Patron']} ({v['Sentido']}, {'en nivel' if v['En nivel'] == 'si' else 'sin nivel'})"
          for v in W["vel"] if v["pos"] >= nw - 2]
    L.append("- Velas S (2 barras): " + ("; ".join(vw) or "ninguna"))
    fg = [f"{tf}: {f['nombre']} — {f['detalle']}" for tf in ("D", "W") for f in r["tf"][tf]["fig"]]
    L.append("- Figuras: " + (" | ".join(fg) or "ninguna"))
    if p:
        L.append(f"- PLAN: entrada {p['e_lo']:.2f}–{p['e_hi']:.2f} ({p['modo']}); stop {p['stop']:.2f} ({p['stop_tipo']}); "
                 f"TP1 {p['tp1']:.2f} (R:R {p['rr1']:.2f}); TP2 {p['tp2']:.2f} (R:R {p['rr2']:.2f}) [{p['tp_tipo']}]; salida principal: cierre semanal < SMA30s (hoy {f2(p['trail'])})")
        L.append(f"- Tamano: riesgo/accion {p['riesgo_acc']:.2f}; por riesgo {cfg['riesgo_max']:.0f}/{p['riesgo_acc']:.2f} = "
                 f"{p['n_riesgo']} acc; por capital {cfg['max_pos']:.0f}/{p['entrada']:.2f} = {p['n_capital']} acc; "
                 f"=> {p['acciones']} acc (limita {p['limite']}); capital {p['capital']:.0f}; riesgo real {p['riesgo_real']:.0f}")
    else:
        extra = f"; nivel a vigilar: cierre sobre {r['activar']:.2f}" if np.isfinite(r["activar"]) else ""
        L.append(f"- Sin plan de compra ({r['razon']}){extra}")
    L.append(f"- Razon: {r['razon']}" + (f" | Alertas: {r['alertas']}" if r["alertas"] else ""))
    nd, nw_ = cfg["barras_ficha"], max(10, cfg["barras_ficha"] * 2 // 3)
    L.append("")
    L.append(f"Diario (ultimas {nd}, fecha,O,H,L,C,Vol):")
    L.append("```")
    for ix, rw in D["d"].iloc[-nd:].iterrows():
        L.append(f"{ix:%Y-%m-%d},{rw['Open']:.2f},{rw['High']:.2f},{rw['Low']:.2f},{rw['Close']:.2f},{int(rw['Volume'])}")
    L.append("```")
    L.append(f"Semanal (ultimas {nw_}):")
    L.append("```")
    for ix, rw in W["d"].iloc[-nw_:].iterrows():
        L.append(f"{ix:%Y-%m-%d},{rw['Open']:.2f},{rw['High']:.2f},{rw['Low']:.2f},{rw['Close']:.2f},{int(rw['Volume'])}")
    L.append("```")
    L.append("")
    return "\n".join(L)


def fila_ranking(r, pos):
    p = r["plan_usable"]
    n = lambda k: round(p[k], 2) if p else np.nan
    return {"Rank": pos, "Ticker": r["ticker"], "Direccion": r["dir"], "Score": r["score"],
            "Cierre": round(r["cierre"], 2), "Ultima barra": r["fecha"].date(),
            "Entrada min": n("e_lo"), "Entrada max": n("e_hi"), "Stop": n("stop"),
            "TP1": n("tp1"), "TP2": n("tp2"), "TP2 tipo": p["tp2_tipo"] if p else "", "Trailing SMA30s sem": n("trail"), "RR TP1": n("rr1"), "RR TP2": n("rr2"),
            "Acciones": p["acciones"] if p else np.nan, "Capital $": round(p["capital"]) if p else np.nan,
            "Riesgo real $": round(p["riesgo_real"]) if p else np.nan,
            "Limite": p["limite"] if p else "", "Modo entrada": p["modo"] if p else "", "Etapa": r["etapa"],
            "RS Mansfield S": round(r["rs_sp"], 2) if np.isfinite(r["rs_sp"]) else np.nan,
            "RS percentil": r.get("rs_pct", np.nan), "Tendencia D": r["tend_d"], "Vigilar (cierre sobre)": round(r["activar"], 2) if np.isfinite(r["activar"]) else np.nan,
            "Razon": r["razon"], "Alertas": r["alertas"]}


def fila_resumen(r):
    D, W = r["tf"]["D"], r["tf"]["W"]
    nd, nw = len(D["d"]), len(W["d"])
    s = lambda x, k=2: round(x, k) if np.isfinite(x) else np.nan
    return {"Ticker": r["ticker"], "Sector": r.get("sector", ""), "Fecha": r["fecha"].date(),
            "Cierre": round(r["cierre"], 2), "Direccion": r["dir"], "Score": r["score"],
            "Etapa sem": r["etapa"], "Sem en etapa": r["etapa_sem"], "Etapa previa": r["etapa_prev"],
            "Pend SMA30s %/4s": s(r["pend_sma30"]), "Dist SMA30s %": s(r["dist_sma30"], 1),
            "Dist max52 %": s(r["dist_max52"], 1), "Tendencia D": r["tend_d"], "Estructura D": r["estructura_txt"],
            "EMA50": round(r["ema50"], 2), "EMA200": round(r["ema200"], 2),
            "RSI dia": round(r["rsi_d"], 1), "RSI sem": round(r["rsi_w"], 1),
            "MACD cruce": r["macd"]["cruce"], "MACD hace (barras)": r["macd"]["hace"] if r["macd"]["cruce"] else np.nan,
            "RS Mansfield S": s(r["rs_sp"]), "RS hace 4s": s(r["rs_sp_4"]), "RS vs sector": s(r["rs_sec"]),
            "Vol ratio dia": s(r["vol_ratio"]), "ATR% dia": s(r["atr_pct"]),
            "Soporte D": s(D["sop"][0][0]) if D["sop"] else np.nan,
            "Resistencia D": s(D["res"][0][0]) if D["res"] else np.nan,
            "Soporte S": s(W["sop"][0][0]) if W["sop"] else np.nan,
            "Resistencia S": s(W["res"][0][0]) if W["res"] else np.nan,
            "POC": s(r["poc"]),
            "Divergencias D": "; ".join(v["tipo"] for v in r["div_d"]),
            "Velas D (5 barras)": "; ".join(sorted({v["Patron"] for v in D["vel"] if v["pos"] >= nd - 5})),
            "Velas S (2 barras)": "; ".join(sorted({v["Patron"] for v in W["vel"] if v["pos"] >= nw - 2})),
            "Figuras D": "; ".join(f["nombre"] for f in D["fig"]),
            "Figuras S": "; ".join(f["nombre"] for f in W["fig"])}


def mensaje_telegram(resultados, previo, fecha, siempre=False):
    comp = {r["ticker"]: r for r in resultados if r["dir"] == "Compra"}
    antes = set((previo or {}).get("compras", []))
    nuevas = [t for t in comp if t not in antes]
    salen = [t for t in antes if t not in comp]
    if not (nuevas or salen or siempre):
        return None
    L = [f"Analisis tecnico · datos al {fecha:%d/%m/%Y}"]
    if nuevas:
        L.append(f"\nNUEVAS COMPRAS ({len(nuevas)}):")
        for t in nuevas:
            p = comp[t]["plan_usable"]
            L.append(f"{t}: entrada {p['e_lo']:.2f}-{p['e_hi']:.2f} | stop {p['stop']:.2f} | TP1 {p['tp1']:.2f} | "
                     f"TP2 {p['tp2']:.2f} | {p['acciones']} acc | {p['capital']:.0f}$ | riesgo {p['riesgo_real']:.0f}$"
                     + (f" | AVISO: {comp[t]['alertas']}" if comp[t]["alertas"] else ""))
    mant = [t for t in comp if t in antes]
    if mant:
        L.append("\nSIGUEN EN COMPRA: " + ", ".join(mant))
    if salen:
        dirs = {r["ticker"]: r["dir"] for r in resultados}
        L.append("\nSALEN DE COMPRA: " + ", ".join(f"{t} ({dirs.get(t, 'sin datos')})" for t in salen))
    if not comp:
        L.append("\nSin compras hoy.")
    return "\n".join(L)


def enviar_telegram(texto, token, chat_id):
    import urllib.parse
    import urllib.request
    trozos, cur = [], ""
    for ln in texto.split("\n"):
        if len(cur) + len(ln) + 1 > 3900:
            trozos.append(cur)
            cur = ""
        cur += ln + "\n"
    trozos.append(cur)
    for tz in trozos:
        data = urllib.parse.urlencode({"chat_id": chat_id, "text": tz}).encode()
        req = urllib.request.Request(f"https://api.telegram.org/bot{token}/sendMessage", data=data)
        urllib.request.urlopen(req, timeout=20).read()


def escribir_excel(ruta, df_rank, df_res, df_vel, df_fig, avisos):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    df_av = pd.DataFrame({"Aviso": avisos or ["(sin avisos)"]})
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        df_rank.to_excel(xw, sheet_name="Ranking", index=False)
        df_res.to_excel(xw, sheet_name="Resumen", index=False)
        df_vel.to_excel(xw, sheet_name="Velas", index=False)
        df_fig.to_excel(xw, sheet_name="Figuras", index=False)
        df_av.to_excel(xw, sheet_name="Avisos", index=False)
    wb = load_workbook(ruta)
    relleno = {"Compra": "C6EFCE", "Venta": "FFC7CE", "Esperar": "FFEB9C"}
    for ws in wb.worksheets:
        ws.freeze_panes = "C2" if ws.title in ("Ranking", "Resumen") else "B2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for col in ws.iter_cols():
            ancho = max(len(str(c.value or "")) for c in col[:200]) + 2
            ws.column_dimensions[col[0].column_letter].width = min(50, max(9, ancho))
        cab = {c.value: c.column for c in ws[1]}
        for nombre in ("Direccion",):
            if nombre in cab:
                for row in ws.iter_rows(min_row=2, min_col=cab[nombre], max_col=cab[nombre]):
                    for cell in row:
                        if cell.value in relleno:
                            cell.fill = PatternFill("solid", fgColor=relleno[cell.value])
        if "Etapa sem" in cab or "Etapa" in cab:
            k = cab.get("Etapa sem", cab.get("Etapa"))
            for row in ws.iter_rows(min_row=2, min_col=k, max_col=k):
                for cell in row:
                    if cell.value == 2:
                        cell.fill = PatternFill("solid", fgColor="C6EFCE")
                    elif cell.value == 4:
                        cell.fill = PatternFill("solid", fgColor="FFC7CE")
    wb.save(ruta)


# ============================================================================
# Principal
# ============================================================================
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tickers", help="Tickers separados por comas (se suman a los del fichero)")
    ap.add_argument("--archivo", default="tickers.txt", help="Fichero de tickers (vacio = S&P 500)")
    ap.add_argument("--salida", default="resultados", help="Carpeta de salida")
    ap.add_argument("--cache", default="cache", help="Carpeta de cache de datos")
    ap.add_argument("--años", type=int, default=6, help="Años de historico en la descarga completa")
    ap.add_argument("--refrescar", action="store_true", help="Fuerza descarga completa")
    ap.add_argument("--solo-cache", action="store_true", help="No usa la red")
    ap.add_argument("--importar-csv", help="Siembra la cache desde un CSV (Ticker,Fecha,Open,...)")
    ap.add_argument("--descartar-ultima", action="store_true", help="Con --importar-csv: quita la ultima fecha (parcial)")
    ap.add_argument("--detalle", choices=["limpio", "normal", "completo"], default="normal")
    ap.add_argument("--top", type=int, default=30, help="Tickers en PDF y fichas cuando se analiza el S&P 500")
    ap.add_argument("--sin-pdf", action="store_true")
    ap.add_argument("--orden", choices=["rs", "score"], default="rs",
                    help="Orden dentro de cada direccion: RS Mansfield (defecto) o score")
    ap.add_argument("--dias", type=int, default=170, help="Velas diarias en el grafico")
    ap.add_argument("--semanas", type=int, default=156, help="Velas semanales en el grafico")
    ap.add_argument("--marcar", type=int, default=15, help="Marcar velas de las ultimas N barras diarias")
    ap.add_argument("--barras-ficha", type=int, default=30, help="Barras diarias OHLCV en fichas.md")
    ap.add_argument("--estocastico", action="store_true")
    ap.add_argument("--telegram", action="store_true", help="Avisa por Telegram si cambian las compras "
                    "(variables TELEGRAM_BOT_TOKEN y TELEGRAM_CHAT_ID)")
    ap.add_argument("--telegram-simular", action="store_true", help="Muestra el mensaje sin enviarlo")
    ap.add_argument("--telegram-siempre", action="store_true", help="Envia aunque no haya cambios")
    ap.add_argument("--cuenta", type=float, default=100000)
    ap.add_argument("--riesgo", type=float, default=0.008, help="Riesgo maximo por operacion (fraccion)")
    ap.add_argument("--max-pos", type=float, default=5000, help="Capital maximo por posicion")
    args = ap.parse_args()

    os.makedirs(args.salida, exist_ok=True)
    cfg = {"dias": args.dias, "semanas": args.semanas, "marcar": args.marcar, "estoc": args.estocastico,
           "riesgo_max": args.cuenta * args.riesgo, "max_pos": args.max_pos,
           "detalle": args.detalle, "barras_ficha": args.barras_ficha}

    if args.importar_csv:
        importar_csv(args.importar_csv, args.cache, args.descartar_ultima)

    # --- Universo
    txt = ""
    if args.archivo and os.path.exists(args.archivo):
        with open(args.archivo, encoding="utf-8") as f:
            txt = f.read()
    tickers = parse_tickers((args.tickers or "") + "," + txt)
    if args.solo_cache:
        p_sp = os.path.join(args.cache, "sp500.csv")
        sp = pd.read_csv(p_sp) if os.path.exists(p_sp) else None
    else:
        sp = cargar_sp500(args.cache)
    if tickers:
        modo = "tickers"
        print(f"Modo: {len(tickers)} tickers de la lista")
    else:
        if sp is None:
            sys.exit("tickers.txt vacio y no se pudo obtener la lista del S&P 500 (sin red ni copia guardada).")
        modo = "sp500"
        tickers = list(dict.fromkeys(sp["Ticker"].tolist()))
        print(f"Modo: S&P 500 completo ({len(tickers)} tickers)")
    sector_de = dict(zip(sp["Ticker"], sp["Sector"])) if sp is not None else {}

    # --- Datos
    pedir = tickers + [BENCH] + ETFS
    datos, avisos = actualizar_cache(pedir, args.cache, args.años, args.refrescar, args.solo_cache)
    sin_etf = [a for a in avisos if a.split(":")[0] in ETFS]
    if sin_etf:
        avisos = [a for a in avisos if a not in sin_etf] + [f"{len(sin_etf)} ETF sectoriales sin datos: RS vs sector no disponible"]
    avisos += descartar_provisional(datos, args.cache)
    ref = max((d.index[-1] for d in datos.values() if len(d)), default=None)
    sin_barras = set()
    for t in tickers:
        if t in datos and len(datos[t]) and ref is not None and (ref - datos[t].index[-1]).days > 4:
            sin_barras.add(t)
            avisos.append(f"{t}: sin barras nuevas desde {datos[t].index[-1]:%Y-%m-%d} "
                          f"(posible delisting, fusion o cambio de ticker); excluido")
    bench_d = datos[BENCH]["Close"] if BENCH in datos else None
    if bench_d is None:
        avisos.append("Sin ^GSPC: no se calculara RS Mansfield")
    bench_w = bench_d.resample("W-FRI").last().dropna() if bench_d is not None else None
    sec_d = {e: datos[e]["Close"] for e in ETFS if e in datos}
    sec_w = {e: s.resample("W-FRI").last().dropna() for e, s in sec_d.items()}

    # --- Analisis
    resultados = []
    for k, t in enumerate(tickers, 1):
        if t not in datos:
            avisos.append(f"{t}: sin datos")
            continue
        if t in sin_barras:
            continue
        if k % 50 == 0:
            print(f"  analizando {k}/{len(tickers)}...")
        etf = SECTOR_ETF.get(sector_de.get(t, ""), None)
        try:
            r, av = analizar(t, datos[t], bench_d, bench_w, sec_d.get(etf), sec_w.get(etf),
                             sector_de.get(t, ""), cfg)
        except Exception as e:  # noqa
            avisos.append(f"{t}: error en el analisis ({type(e).__name__}: {e})")
            continue
        if r is None:
            avisos.append(f"{t}: {av}")
        else:
            resultados.append(r)
            avisos += av
    if not resultados:
        sys.exit("Sin resultados.")

    orden = {"Compra": 0, "Esperar": 1, "Venta": 2}
    rs_s = pd.Series({r["ticker"]: r["rs_sp"] for r in resultados})
    pct_rs = (rs_s.rank(pct=True) * 100).round(0)
    for r in resultados:
        r["rs_pct"] = pct_rs.get(r["ticker"], np.nan)
    if args.orden == "rs":
        # En el backtest, ordenar por RS cuando hay mas señales que caja fue lo que mas rindio
        resultados.sort(key=lambda r: (orden[r["dir"]], -(r["rs_sp"] if np.isfinite(r["rs_sp"]) else -1e9), -r["score"]))
    else:
        resultados.sort(key=lambda r: (orden[r["dir"]], -r["score"]))
    sel = resultados if modo == "tickers" else resultados[:args.top]

    # --- Ranking, Excel, fichas
    df_rank = pd.DataFrame([fila_ranking(r, i) for i, r in enumerate(resultados, 1)])
    df_res = pd.DataFrame([fila_resumen(r) for r in resultados])
    velas = [{k: v for k, v in x.items() if k != "pos"} for r in resultados
             for x in r["tf"]["D"]["vel"] if x["pos"] >= len(r["dd"]) - 5]
    velas += [{k: v for k, v in x.items() if k != "pos"} for r in resultados
              for x in r["tf"]["W"]["vel"] if x["pos"] >= len(r["tf"]["W"]["d"]) - 2]
    df_vel = pd.DataFrame(velas, columns=["Ticker", "TF", "Fecha", "Patron", "Sentido", "En nivel"])
    df_fig = pd.DataFrame([{"Ticker": r["ticker"], "TF": tf, "Figura": f["nombre"], "Detalle": f["detalle"]}
                           for r in resultados for tf in ("D", "W") for f in r["tf"][tf]["fig"]],
                          columns=["Ticker", "TF", "Figura", "Detalle"])
    df_rank.to_csv(os.path.join(args.salida, "ranking.csv"), index=False, encoding="utf-8-sig")
    try:
        escribir_excel(os.path.join(args.salida, "resumen_patrones.xlsx"), df_rank, df_res, df_vel, df_fig, avisos)
    except ImportError:
        df_res.to_csv(os.path.join(args.salida, "resumen.csv"), index=False)
    with open(os.path.join(args.salida, "fichas.md"), "w", encoding="utf-8") as f:
        f.write(f"# Fichas tecnicas · {resultados[0]['fecha']:%Y-%m-%d} · cuenta {args.cuenta:.0f} · "
                f"riesgo {cfg['riesgo_max']:.0f} · max posicion {args.max_pos:.0f}\n\n")
        f.write("Resumen: " + ", ".join(f"{d}: {sum(1 for r in resultados if r['dir'] == d)}"
                                         for d in ("Compra", "Esperar", "Venta")) + "\n\n")
        for r in sel:
            f.write(ficha(r, cfg) + "\n")
    with open(os.path.join(args.salida, "avisos.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(avisos) or "(sin avisos)")
    compras = [r["ticker"] for r in resultados if r["dir"] == "Compra"]
    with open(os.path.join(args.salida, "compras.txt"), "w", encoding="utf-8") as f:
        f.write(",".join(t.replace("-", ".") for t in compras))
    filas = []
    for t in tickers:
        d = datos.get(t)
        if d is None or not len(d):
            filas.append({"Ticker": t, "Primera": "", "Ultima": "", "Barras": 0, "Dias atraso": np.nan,
                          "Estado": "sin datos"})
            continue
        est = "sin barras nuevas" if t in sin_barras else "historia corta" if len(d) < MIN_BARRAS_D else "ok"
        filas.append({"Ticker": t, "Primera": d.index[0].date(), "Ultima": d.index[-1].date(),
                      "Barras": len(d), "Dias atraso": (ref - d.index[-1]).days, "Estado": est})
    pd.DataFrame(filas).to_csv(os.path.join(args.salida, "estado_datos.csv"), index=False)

    # --- PDF
    if not args.sin_pdf:
        with PdfPages(os.path.join(args.salida, "graficos_tecnicos.pdf")) as pdf:
            for i, r in enumerate(sel, 1):
                print(f"Grafico {i}/{len(sel)}: {r['ticker']}")
                for tf in ("D", "W"):
                    dibujar(pdf, r, tf, cfg)

    if args.telegram or args.telegram_simular:
        ruta_est = os.path.join(args.cache, "estado_compras.json")
        previo = None
        if os.path.exists(ruta_est):
            try:
                previo = json.load(open(ruta_est, encoding="utf-8"))
            except Exception:
                previo = None
        msg = mensaje_telegram(resultados, previo, resultados[0]["fecha"], args.telegram_siempre)
        guardar = False
        if msg is None:
            print("Telegram: sin cambios en las compras; no se envia")
            guardar = True
        elif args.telegram_simular:
            print("--- Telegram (simulado) ---\n" + msg + "\n---")
        else:
            tok, chat = os.environ.get("TELEGRAM_BOT_TOKEN"), os.environ.get("TELEGRAM_CHAT_ID")
            if not tok or not chat:
                print("Telegram: faltan TELEGRAM_BOT_TOKEN / TELEGRAM_CHAT_ID; no se envia")
            else:
                try:
                    enviar_telegram(msg, tok, chat)
                    print("Telegram: mensaje enviado")
                    guardar = True
                except Exception as e:  # noqa
                    print(f"Telegram: error al enviar ({type(e).__name__}); se reintentara en la proxima ejecucion")
        if guardar and not args.telegram_simular:
            os.makedirs(args.cache, exist_ok=True)
            with open(ruta_est, "w", encoding="utf-8") as f:
                json.dump({"fecha": str(resultados[0]["fecha"].date()), "compras": compras}, f)

    print(f"\nListo. {len(resultados)} analizados, {len(sel)} en PDF/fichas. Resultados en: {os.path.abspath(args.salida)}")
    print("Ranking: " + ", ".join(f"{r['ticker']}({r['dir'][0]}{r['score']})" for r in resultados[:10]))
    if avisos:
        print(f"{len(avisos)} avisos (ver avisos.txt)")


if __name__ == "__main__":
    main()
